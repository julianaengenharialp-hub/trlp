import io, json, re, zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
import streamlit as st
from docx import Document
from docx.shared import Cm, Pt

st.set_page_config(page_title="TRLP | Elaboração orientada", layout="wide")
TITULOS = ["OBJETO","JUSTIFICATIVA E OBJETIVO DA CONTRATAÇÃO","DESCRIÇÃO DA SOLUÇÃO COMO UM TODO","REQUISITOS DA CONTRATAÇÃO","FORMA DE EXECUÇÃO DO OBJETO","GESTÃO DO CONTRATO","MEDIÇÃO E PAGAMENTO","SELEÇÃO DO FORNECEDOR","ESTIMATIVA DO PREÇO","ADEQUAÇÃO ORÇAMENTÁRIA","DISPOSIÇÕES FINAIS"]
ARQUIVO = Path(__file__).with_name("modelo_tr.odt")
T = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
O = "{urn:oasis:names:tc:opendocument:xmlns:office:1.0}"
NUM = re.compile(r"^(\d+(?:\.\d+)+)\.?\s+(.+)$")
CAP = re.compile(r"^(\d{1,2})\.\s+([A-ZÁÉÍÓÚÃÕÇÊÔÜ0-9 /–-]+)$")
EXEMPLOS = ("campo de futebol","quadra de basquete","pista de caminhada","parquinho infantil","jardim do caju","secretaria de esportes","projeto padrão federal")
INSTRUCOES = ("observação:","o texto abaixo","texto abaixo","modelo 1:","modelo 2:","modelo 3:","caso seja autorizado","caso não tenha memorial","segue abaixo exemplo","analisar cada caso","aqui é um exemplo","deve ser analisado cada caso","se houver outras normas","descrever detalhadamente","caso tenha havido o etp")
def limpar(s): return " ".join(s.split())
@st.cache_data
def carregar(b):
    with zipfile.ZipFile(io.BytesIO(b)) as z:
        root = ET.fromstring(z.read("content.xml"))
    body = root.find(f".//{O}body/{O}text")
    def texto(node):
        if node.tag in (O+"annotation",O+"annotation-end"): return ""
        s=node.text or ""
        for child in node: s+=texto(child)+(child.tail or "")
        return s
    return [limpar(texto(p)) for p in body.iter() if p.tag in (T+"p",T+"h") and limpar(texto(p))]
def categorizar(linhas):
    caps={n:[] for n in range(1,12)}
    cap=None; grupo=None
    for s in linhas:
        m=CAP.match(s)
        if m and int(m.group(1)) in caps and m.group(2).strip() in TITULOS:
            cap=int(m.group(1));grupo=None;continue
        if cap is None:continue
        low=s.lower()
        if any(x in low for x in INSTRUCOES): 
            if low.startswith("modelo "):grupo=low
            continue
        n=NUM.match(s)
        if n:
            numero,conteudo=n.groups()
            if int(numero.split(".")[0])!=cap:continue
            caps[cap].append({"id":f"{cap}:{len(caps[cap])}","numero":numero,"texto":conteudo,"grupo":grupo})
        elif caps[cap] and s and not any(x in low for x in EXEMPLOS):
            # Continuação não numerada: mantida somente quando for um texto claramente integrante do item.
            if s[0].islower() or s.startswith(("a)","b)","c)","d)")):
                caps[cap][-1]["texto"]+=" "+s
    return caps
def suspeito(item):
    t=item["texto"].lower()
    return any(x in t for x in EXEMPLOS) or "________________" in t or "[descrever" in t or "120 (cento" in t or "educ.suprimentos@" in t or "3269-7800" in t or "contados da data das respectivas medições" in t
def inicial(item):
    return "Revisar / adaptar" if suspeito(item) else "Manter texto do modelo"
def get(k,default=""):
    return st.session_state.get(k,default)
def campo(label,key,valor=""):
    return st.text_input(label,key=key,value=valor) if key not in st.session_state else st.text_input(label,key=key)
def lista(label,key,opcoes):
    return st.selectbox(label,opcoes,key=key)
def escolha(label,key):
    return lista(label,key,["Selecione","Sim","Não"])
def registro():
    return {k:v for k,v in st.session_state.items() if k.startswith(("cfg_","opt_","edit_","why_"))}
def trocar_numero(s,numero):
    return f"{numero}. {s.strip()}"
st.title("TRLP | Termos de Referência")
st.caption("Versão 0.4 experimental | Cláusulas reutilizáveis, opções de aplicação e edição pontual. Revisão técnica e jurídica permanece necessária.")
if not ARQUIVO.exists():
    st.error("Coloque o arquivo modelo_tr.odt na raiz do repositório.");st.stop()
linhas=carregar(ARQUIVO.read_bytes())
caps=categorizar(linhas)
st.sidebar.success("Modelo ODT carregado")
st.sidebar.caption("Os exemplos do modelo não são presumidos obrigatórios. Nenhum dado é salvo no GitHub.")
up=st.sidebar.file_uploader("Retomar trabalho (.json)",type="json")
if up and st.sidebar.button("Importar"):
    try:
        dados=json.load(up)
        if dados.get("formato")!="trlp-04":raise ValueError("Formato de arquivo incompatível.")
        for k,v in dados["campos"].items():st.session_state[k]=v
        st.rerun()
    except Exception as e:st.sidebar.error(str(e))
a,b,c=st.tabs(["1. Configuração","2. Cláusulas por capítulo","3. Conferência e exportação"])
with a:
    x,y=st.columns(2)
    with x:
        campo("Unidade solicitante","cfg_unidade","Secretaria de Obras e Infraestrutura")
        campo("Objeto da contratação (resumo indispensável)","cfg_objeto")
        lista("Natureza","cfg_natureza",["Selecione","Obra de engenharia","Serviço de engenharia","Outro serviço"])
        lista("Regime de execução","cfg_regime",["Selecione","Empreitada por preço global","Empreitada por preço unitário","Tarefa","Outro / análise necessária"])
    with y:
        escolha("Foi elaborado ETP?","cfg_etp")
        escolha("Existe memorial descritivo anexo?","cfg_memorial")
        escolha("Será admitida subcontratação?","cfg_sub")
        escolha("Será exigida garantia contratual?","cfg_garantia")
        escolha("Será prevista visita técnica?","cfg_visita")
        campo("Prazo de execução (dias)","cfg_prazo")
    st.info("Nenhuma escolha de natureza técnica é presumida. Preencha somente os dados específicos; as cláusulas são tratadas na aba seguinte.")
with b:
    sec=st.selectbox("Capítulo",range(1,12),format_func=lambda i:f"{i}. {TITULOS[i-1]}")
    st.subheader(f"{sec}. {TITULOS[sec-1]}")
    st.caption("Manter = usar a redação sugerida; adaptar = editar somente esta cláusula; não aplicar = registrar para conferência. Itens do exemplo exigem revisão.")
    itens=caps[sec]
    if not itens:st.warning("Nenhum subitem numerado foi identificado automaticamente neste capítulo. É necessária análise manual.")
    for item in itens:
        num=item["numero"];ident=item["id"];texto=item["texto"]
        if sec==5 and num.startswith("5.1.") and item["grupo"] and get("cfg_regime")!="Selecione":
            g=item["grupo"]
            escolhido=get("cfg_regime").lower()
            if ("global" in g and "global" not in escolhido) or ("unitário" in g and "unitário" not in escolhido) or ("tarefa" in g and "tarefa" not in escolhido):
                continue
        if num=="4.10.1" and get("cfg_garantia")=="Não":continue
        if num.startswith("4.11.") and get("cfg_visita")=="Não":continue
        with st.expander(f"{num}. {texto[:90]}{'…' if len(texto)>90 else ''}"):
            st.caption("Texto extraído do modelo municipal:")
            st.write(texto)
            if suspeito(item):st.warning("Contém exemplo, campo variável ou condição que exige análise. Não reutilizar sem conferir.")
            key="opt_"+ident
            opcoes=["Manter texto do modelo","Revisar / adaptar","Não aplicar (justificar)"]
            if key not in st.session_state:st.session_state[key]=inicial(item)
            modo=st.radio("Tratamento",opcoes,key=key,horizontal=True)
            if modo=="Revisar / adaptar":
                ek="edit_"+ident
                if ek not in st.session_state:st.session_state[ek]=texto
                st.text_area("Redação desta cláusula",key=ek,height=125)
                st.text_input("Motivo da alteração (para relatório interno)",key="why_"+ident)
            elif modo.startswith("Não aplicar"):
                st.text_input("Motivo da não aplicação (para relatório interno)",key="why_"+ident)
    st.divider()
    st.text_area("Complemento técnico do capítulo (somente se necessário)",key=f"cfg_extra_{sec}",height=90)
def montar():
    secoes={};alertas=[];mudancas=[]
    for n in range(1,12):
        trechos=[]
        for item in caps[n]:
            numero=item["numero"];ident=item["id"]
            if n==5 and numero.startswith("5.1.") and item["grupo"]:
                regime=get("cfg_regime").lower();g=item["grupo"]
                if regime=="selecione" or not regime:continue
                if not (("global" in regime and "global" in g) or ("unitário" in regime and "unitário" in g) or ("tarefa" in regime and "tarefa" in g)):continue
            if numero=="4.10.1" and get("cfg_garantia")=="Não":continue
            if numero.startswith("4.11.") and get("cfg_visita")=="Não":continue
            modo=get("opt_"+ident,inicial(item))
            if modo.startswith("Não aplicar"):
                mudancas.append((numero,"Não aplicado",get("why_"+ident)))
                if not get("why_"+ident).strip():alertas.append(f"{numero}: justificar não aplicação")
                continue
            txt=get("edit_"+ident,item["texto"]) if modo=="Revisar / adaptar" else item["texto"]
            if modo=="Revisar / adaptar":
                mudancas.append((numero,"Adaptado",get("why_"+ident)))
                if not get("why_"+ident).strip():alertas.append(f"{numero}: justificar alteração")
            if suspeito(item) and modo=="Manter texto do modelo":alertas.append(f"{numero}: possível texto exemplificativo mantido")
            if not txt.strip():alertas.append(f"{numero}: texto vazio");continue
            trechos.append(trocar_numero(txt,numero))
        if get(f"cfg_extra_{n}").strip():trechos.append("COMPLEMENTO TÉCNICO (revisar posicionamento e numeração): "+get(f"cfg_extra_{n}").strip())
        secoes[n]=trechos
    return secoes,alertas,mudancas
with c:
    secoes,alertas,mudancas=montar()
    for chave,rotulo in [("cfg_unidade","Unidade solicitante"),("cfg_objeto","Objeto"),("cfg_natureza","Natureza"),("cfg_regime","Regime de execução"),("cfg_etp","ETP"),("cfg_memorial","Memorial"),("cfg_sub","Subcontratação"),("cfg_garantia","Garantia"),("cfg_visita","Visita técnica")]:
        if get(chave) in ("","Selecione"):alertas.append(f"Informar: {rotulo}")
    if get("cfg_sub")=="Sim":alertas.append("Conferir e delimitar parcelas subcontratáveis nos itens 4.9 e 5.5")
    if get("cfg_prazo"):alertas.append("Conferir prazo de execução no item 5.2: o exemplo pode conter prazo diferente")
    if alertas:
        st.warning(f"{len(alertas)} pontos para conferir")
        with st.expander("Ver pendências e alertas",expanded=True):
            for x in dict.fromkeys(alertas):st.write("• "+x)
    else:st.success("Nenhum alerta automático identificado. Isso não equivale a validação jurídica.")
    st.caption(f"Cláusulas alteradas ou não aplicadas: {len(mudancas)}")
    for numero,tipo,motivo in mudancas:st.write(f"**{numero}** · {tipo} · {motivo or 'Justificativa pendente'}")
    for n,trechos in secoes.items():
        with st.expander(f"{n}. {TITULOS[n-1]} ({len(trechos)} trechos)",expanded=n==1):
            st.text("\n\n".join(trechos) or "[CAPÍTULO SEM TEXTO GERADO]")
    dados={"formato":"trlp-04","campos":registro()}
    st.download_button("Salvar trabalho (.json)",json.dumps(dados,ensure_ascii=False,indent=2).encode(),"trabalho_trlp_v04.json","application/json")
    doc=Document()
    doc.sections[0].top_margin=Cm(2.3);doc.sections[0].bottom_margin=Cm(2.3)
    doc.add_heading("TERMO DE REFERÊNCIA | RASCUNHO PARA CONFERÊNCIA",0)
    doc.add_paragraph("UNIDADE SOLICITANTE: "+get("cfg_unidade"))
    doc.add_paragraph("OBJETO INFORMADO: "+get("cfg_objeto"))
    for n,trechos in secoes.items():
        doc.add_heading(f"{n}. {TITULOS[n-1]}",level=1)
        for trecho in trechos:doc.add_paragraph(trecho)
        if not trechos:doc.add_paragraph("[CAPÍTULO A COMPLETAR]")
    doc.add_page_break();doc.add_heading("RELATÓRIO INTERNO DE CONFERÊNCIA",1)
    for x in dict.fromkeys(alertas):doc.add_paragraph("PENDÊNCIA: "+x)
    for numero,tipo,motivo in mudancas:doc.add_paragraph(f"{numero}: {tipo}; justificativa: {motivo or 'não informada'}")
    doc.add_paragraph("RASCUNHO EXPERIMENTAL: conferir integralmente contra o modelo municipal antes de encaminhar à licitação.")
    out=io.BytesIO();doc.save(out)
    st.download_button("Baixar rascunho DOCX",out.getvalue(),"TRLP_rascunho_v04.docx","application/vnd.openxmlformats-officedocument.wordprocessingml.document")
st.sidebar.warning("Protótipo: a extração automática ainda não distingue todas as instruções, exemplos, listas e alternativas. Não usar o DOCX diretamente em licitação.")
